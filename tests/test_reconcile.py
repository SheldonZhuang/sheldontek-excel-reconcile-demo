import sys
from datetime import date
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import make_samples  # noqa: E402
import reconcile as rc  # noqa: E402


@pytest.fixture(scope="session")
def samples(tmp_path_factory):
    return make_samples.build(tmp_path_factory.mktemp("data"))


def run(samples, fuzzy=None, kind="xlsx"):
    a, _ = rc.load(str(samples[f"a_{kind}"]))
    b, _ = rc.load(str(samples[f"b_{kind}"]))
    return a, b, rc.reconcile(a, b, "id", fuzzy)


def keys(df, idx):
    return sorted(str(df.iloc[i]["id"]) for i in idx)


# ---- normalisation rules
@pytest.mark.parametrize("raw,expected", [
    ("  Sample   Co 004 ", "sample co 004"), ("SAMPLE CO", "sample co"), (None, ""),
    (float("nan"), ""), ("a\tb\nc", "a b c"),
])
def test_norm_text(raw, expected):
    assert rc.norm_text(raw) == expected


@pytest.mark.parametrize("raw", ["Sample Co 004 Limited", "  sample co 004   ltd ",
                                 "Sample Co. 004 Ltd.", "SAMPLE CO 004 INC"])
def test_norm_name_drops_suffix_and_punctuation(raw):
    assert rc.norm_name(raw) == "sample co 004"


@pytest.mark.parametrize("raw,expected", [
    ("1,250.00", 1250.0), (" 1250 ", 1250.0), (1250, 1250.0), (1250.0, 1250.0),
    ("-3.5", -3.5), ("12 apples", None), ("", None), (True, None),
])
def test_norm_number(raw, expected):
    assert rc.norm_number(raw) == expected


def test_norm_date_formats_and_dayfirst():
    for raw in ("2026-02-03", "03/02/2026", "03.02.2026", "3 Feb 2026", "Feb 3, 2026",
                pd.Timestamp("2026-02-03"), date(2026, 2, 3)):
        assert rc.norm_date(raw) == date(2026, 2, 3), raw
    assert rc.norm_date("03/02/2026", dayfirst=False) == date(2026, 3, 2)
    assert rc.norm_date("not a date") is None


def test_norm_key():
    assert rc.norm_key(1001) == rc.norm_key("1001") == rc.norm_key(1001.0) == rc.norm_key(" 1001 ")
    assert rc.norm_key("1O26") != rc.norm_key("1026")


def test_same_numbers_guard():
    assert rc.same_numbers("sample co 017", "sample co 17")
    assert not rc.same_numbers("sample co 017", "sample co 071")


# ---- matching
def test_exact_key_only(samples):
    a, b, res = run(samples)
    assert len(res.matched) == 26 and all(m["how"] == "key" for m in res.matched)
    assert keys(a, res.only_a) == ["1005", "1017", "1026", "1029"]
    assert keys(b, res.only_b) == ["1071", "1101", "1102", "1O26"]


def test_fuzzy_pairs_mistyped_ids(samples):
    a, b, res = run(samples, ("name", 90))
    fuzzy = {(str(a.iloc[m["a"]]["id"]), str(b.iloc[m["b"]]["id"])) for m in res.matched
             if m["how"] != "key"}
    assert fuzzy == {("1017", "1071"), ("1026", "1O26")}
    assert keys(a, res.only_a) == ["1005", "1029"]
    assert keys(b, res.only_b) == ["1101", "1102"]


def test_fuzzy_decoy_with_different_number_not_matched(samples):
    a, b, res = run(samples, ("name", 50))   # even with a very low threshold
    b_names = {b.iloc[m["b"]]["name"] for m in res.matched}
    assert "Sample Co 071 Limited" not in b_names


def test_fuzzy_threshold_respected():
    a = pd.DataFrame({"id": [1], "name": ["Sample Co 001 Northern Trading"]})
    b = pd.DataFrame({"id": [9], "name": ["Sample Co 001 Trading"]})
    s = rc.fuzz.token_sort_ratio(rc.norm_name(a.name[0]), rc.norm_name(b.name[0]))
    assert len(rc.reconcile(a, b, "id", ("name", int(s) + 1)).matched) == 0
    assert len(rc.reconcile(a, b, "id", ("name", int(s))).matched) == 1


def test_fuzzy_is_one_to_one():
    a = pd.DataFrame({"id": [1, 2], "name": ["Sample Co 001", "Sample Co 001 Ltd"]})
    b = pd.DataFrame({"id": [9], "name": ["sample co 001"]})
    res = rc.reconcile(a, b, "id", ("name", 90))
    assert len(res.matched) == 1 and len(res.only_a) == 1


def test_duplicate_keys_reported():
    a = pd.DataFrame({"id": [1, 1, 2], "v": [1, 2, 3]})
    b = pd.DataFrame({"id": [1, 2], "v": [1, 3]})
    res = rc.reconcile(a, b, "id")
    assert res.duplicates == {"A": ["1"]} and len(res.only_a) == 1


def test_missing_key_column_raises():
    with pytest.raises(KeyError):
        rc.reconcile(pd.DataFrame({"x": [1]}), pd.DataFrame({"id": [1]}), "id")


# ---- diff detection
def diff_set(a, res):
    return {(str(a.iloc[res.matched[d["pair"]]["a"]]["id"]), d["column"]) for d in res.diffs}


EXPECTED = {("1003", "amount"), ("1011", "amount"), ("1022", "amount"), ("1008", "status"),
            ("1013", "status"), ("1019", "invoice_date"), ("1017", "id"), ("1026", "id")}


def test_diffs_exactly_the_planted_ones(samples):
    a, b, res = run(samples, ("name", 90))
    assert diff_set(a, res) == EXPECTED


def test_cosmetic_noise_is_not_a_difference(samples):
    # rows 1004 (name case/space/Ltd), 1010 (amount "1,250.00"-style text),
    # 1006 (d/m/y date text), 1007 (status upper-case + space) must be clean
    a, b, res = run(samples, ("name", 90))
    noisy = {"1004", "1006", "1007", "1010", "1012", "1014", "1015", "1020", "1021", "1028"}
    assert not {k for k, _ in diff_set(a, res)} & noisy


def test_csv_and_two_sheet_inputs_agree(samples):
    a, _, x = run(samples, ("name", 90), "xlsx")
    ac, _, c = run(samples, ("name", 90), "csv")
    a2, _ = rc.load(f"{samples['two_sheets']}#System A")
    b2, _ = rc.load(f"{samples['two_sheets']}#System B")
    t = rc.reconcile(a2, b2, "id", ("name", 90))
    assert diff_set(a, x) == diff_set(ac, c) == diff_set(a2, t) == EXPECTED


# ---- report
def test_report_sheets_and_highlights(samples, tmp_path):
    out = tmp_path / "report.xlsx"
    assert rc.main([str(samples["a_xlsx"]), str(samples["b_xlsx"]), "--key", "id",
                    "--fuzzy", "name:90", "-o", str(out)]) == 0
    wb = load_workbook(out)
    assert wb.sheetnames == ["Summary", "Matched", "Only in A", "Only in B", "Value differences"]
    summary = {r[0]: r[1] for r in wb["Summary"].iter_rows(values_only=True) if r[0]}
    assert summary["Matched pairs"] == 28 and summary["Differing cells"] == 8
    assert summary["Only in A"] == 2 and summary["Only in B"] == 2
    ws = wb["Matched"]
    head = [c.value for c in ws[1]]
    amt_a, amt_b = head.index("amount (A)") + 1, head.index("amount (B)") + 1
    row = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, head.index("id (A)") + 1).value == 1003)
    assert ws.cell(row, amt_a).fill.fgColor.rgb.endswith("FFE08A")
    assert ws.cell(row, amt_b).fill.fgColor.rgb.endswith("FFE08A")
    clean = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, head.index("id (A)") + 1).value == 1002)
    assert ws.cell(clean, amt_a).fill.fill_type is None
    assert wb["Value differences"].max_row == 1 + 8


def test_cli_fuzzy_argument_validation():
    with pytest.raises(SystemExit):
        rc.main(["a.csv", "b.csv", "--key", "id", "--fuzzy", "name"])
