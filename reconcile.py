"""Reconcile two datasets by key and write an Excel report.

Usage:
    python reconcile.py a.xlsx b.xlsx --key id [--fuzzy name:90] [-o report.xlsx]
    python reconcile.py a.csv  b.csv  --key id
    python reconcile.py book.xlsx#"System A" book.xlsx#"System B" --key id

Steps
  1. load both sides (CSV, or an Excel sheet; "file.xlsx#Sheet" picks a sheet)
  2. normalise every compared value (trim, collapse spaces, case, numbers, dates;
     names also drop company suffixes and punctuation)
  3. exact match on the normalised key
  4. optional fuzzy match on a name column for rows still unmatched
     (rapidfuzz token_sort_ratio >= threshold, numbers in both names must agree,
     each row is used at most once, best score first)
  5. compare the other shared columns of every matched pair
  6. report: Summary, Matched, Only in A, Only in B, Value differences
"""
from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from rapidfuzz import fuzz

COMPANY_SUFFIXES = {"ltd", "limited", "inc", "llc", "co", "corp", "plc", "gmbh", "company"}
_WS = re.compile(r"\s+")
_NUM_TEXT = re.compile(r"^[-+]?[\d,]*\.?\d+$")
_DIGITS = re.compile(r"\d+")
DATE_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y/%m/%d",
                "%d %b %Y", "%d %B %Y", "%b %d, %Y", "%B %d, %Y")
DATE_FORMATS_MDY = ("%Y-%m-%d", "%m/%d/%Y", "%m-%d-%Y", "%Y/%m/%d",
                    "%d %b %Y", "%d %B %Y", "%b %d, %Y", "%B %d, %Y")


# ------------------------------------------------------------------ normalise
def norm_text(v) -> str:
    """Trim, collapse inner whitespace, casefold. None/NaN -> ''."""
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    return _WS.sub(" ", str(v)).strip().casefold()


def norm_name(v) -> str:
    """Text rules plus: drop punctuation and trailing company suffixes."""
    words = re.sub(r"[^\w\s]", " ", norm_text(v)).split()
    while words and words[-1] in COMPANY_SUFFIXES:
        words.pop()
    return " ".join(words)


def norm_number(v):
    """1,250.00 / ' 1250 ' / 1250.0 -> 1250.0 ; otherwise None."""
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)) and not pd.isna(v):
        return round(float(v), 6)
    t = norm_text(v).replace(" ", "")
    if t and _NUM_TEXT.match(t):
        return round(float(t.replace(",", "")), 6)
    return None


def norm_date(v, dayfirst: bool = True):
    """datetime / date / text in common formats -> date ; otherwise None.
    Ambiguous d/m vs m/d text follows `dayfirst` (default: day first)."""
    if isinstance(v, (pd.Timestamp, datetime)):
        return v.date()
    if isinstance(v, date):
        return v
    t = _WS.sub(" ", str(v)).strip() if v is not None else ""
    for fmt in (DATE_FORMATS if dayfirst else DATE_FORMATS_MDY):
        try:
            return datetime.strptime(t, fmt).date()
        except ValueError:
            continue
    return None


def norm_key(v) -> str:
    """Keys: text rules, and 1001 / 1001.0 / ' 1001' all become '1001'."""
    n = norm_number(v)
    if n is not None and float(n).is_integer():
        return str(int(n))
    return norm_text(v)


def norm_value(v, dayfirst: bool = True):
    """Comparable form of a cell: number, then date, then normalised text."""
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    n = norm_number(v)
    if n is not None:
        return n
    d = norm_date(v, dayfirst)
    if d is not None:
        return d
    return norm_text(v)


def same_numbers(a: str, b: str) -> bool:
    """'sample co 017' vs 'sample co 071' -> False: the digits must agree."""
    return [int(x) for x in _DIGITS.findall(a)] == [int(x) for x in _DIGITS.findall(b)]


# ------------------------------------------------------------------ load
def load(spec: str) -> tuple[pd.DataFrame, str]:
    """'file.csv', 'file.xlsx' (first sheet) or 'file.xlsx#Sheet name'."""
    path, _, sheet = spec.partition("#")
    p = Path(path)
    if p.suffix.lower() == ".csv":
        df = pd.read_csv(p, dtype=object, keep_default_na=False)
        label = p.name
    else:
        df = pd.read_excel(p, sheet_name=sheet or 0, dtype=object)
        label = f"{p.name}#{sheet}" if sheet else p.name
    df.columns = [str(c).strip() for c in df.columns]
    return df, label


# ------------------------------------------------------------------ reconcile
@dataclass
class Result:
    key: str
    columns: list[str]                       # compared columns (shared, minus key)
    matched: list[dict] = field(default_factory=list)   # {"a": i, "b": j, "how": ..., "score": ...}
    only_a: list[int] = field(default_factory=list)
    only_b: list[int] = field(default_factory=list)
    diffs: list[dict] = field(default_factory=list)     # {"pair": k, "column": c, "a": .., "b": ..}
    duplicates: dict = field(default_factory=dict)       # side -> list of duplicated keys


def reconcile(a: pd.DataFrame, b: pd.DataFrame, key: str,
              fuzzy: tuple[str, int] | None = None,
              name_cols: list[str] | None = None,
              dayfirst: bool = True) -> Result:
    for side, df in (("A", a), ("B", b)):
        if key not in df.columns:
            raise KeyError(f"key column '{key}' not in dataset {side}: {list(df.columns)}")
    if fuzzy and (fuzzy[0] not in a.columns or fuzzy[0] not in b.columns):
        raise KeyError(f"fuzzy column '{fuzzy[0]}' must exist in both datasets")
    columns = [c for c in a.columns if c in b.columns and c != key]
    if name_cols is None:
        name_cols = [fuzzy[0]] if fuzzy else [c for c in columns if "name" in c.lower()]
    res = Result(key, columns)

    ka = [norm_key(v) for v in a[key]]
    kb = [norm_key(v) for v in b[key]]
    for side, keys in (("A", ka), ("B", kb)):
        seen, dup = set(), []
        for k in keys:
            if k in seen and k not in dup:
                dup.append(k)
            seen.add(k)
        if dup:
            res.duplicates[side] = dup

    # 1) exact key match (first occurrence of a key on each side)
    index_b: dict[str, int] = {}
    for j, k in enumerate(kb):
        index_b.setdefault(k, j)
    used_b: set[int] = set()
    left_a: list[int] = []
    for i, k in enumerate(ka):
        j = index_b.get(k)
        if k and j is not None and j not in used_b:
            res.matched.append({"a": i, "b": j, "how": "key", "score": 100})
            used_b.add(j)
        else:
            left_a.append(i)
    left_b = [j for j in range(len(b)) if j not in used_b]

    # 2) fuzzy on name for what is left, best pairs first, one-to-one
    if fuzzy and left_a and left_b:
        col, threshold = fuzzy
        na = {i: norm_name(a.at[a.index[i], col]) for i in left_a}
        nb = {j: norm_name(b.at[b.index[j], col]) for j in left_b}
        cands = []
        for i in left_a:
            for j in left_b:
                if not na[i] or not nb[j] or not same_numbers(na[i], nb[j]):
                    continue
                s = fuzz.token_sort_ratio(na[i], nb[j])
                if s >= threshold:
                    cands.append((s, i, j))
        cands.sort(key=lambda t: (-t[0], t[1], t[2]))
        taken_a, taken_b = set(), set()
        for s, i, j in cands:
            if i in taken_a or j in taken_b:
                continue
            res.matched.append({"a": i, "b": j, "how": f"fuzzy {col}", "score": round(s, 1)})
            taken_a.add(i)
            taken_b.add(j)
        left_a = [i for i in left_a if i not in taken_a]
        left_b = [j for j in left_b if j not in taken_b]

    res.only_a, res.only_b = left_a, left_b
    res.matched.sort(key=lambda m: m["a"])

    # 3) value differences on matched pairs
    for k, m in enumerate(res.matched):
        ra, rb = a.iloc[m["a"]], b.iloc[m["b"]]
        if m["how"] != "key" and norm_key(ra[key]) != norm_key(rb[key]):
            res.diffs.append({"pair": k, "column": key, "a": ra[key], "b": rb[key]})
        for c in columns:
            f = norm_name if c in name_cols else (lambda v: norm_value(v, dayfirst))
            if f(ra[c]) != f(rb[c]):
                res.diffs.append({"pair": k, "column": c, "a": ra[c], "b": rb[c]})
    return res


# ------------------------------------------------------------------ report
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
HEAD_FONT = Font(bold=True, color="FFFFFF")
DIFF_FILL = PatternFill("solid", fgColor="FFE08A")   # amber: value differs
FUZZY_FILL = PatternFill("solid", fgColor="D6E9FF")  # blue: paired by fuzzy name
ONLY_FILL = PatternFill("solid", fgColor="F8D7DA")


def _cell(v):
    if isinstance(v, float) and pd.isna(v):
        return None
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime()
    return v


def _sheet(wb: Workbook, title: str, header: list[str], rows: list[list]):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
    for r in rows:
        ws.append([_cell(v) for v in r])
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(header, 1):
        width = max([len(str(h))] + [len(str(_cell(r[i - 1]) or "")) for r in rows])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 40)
        for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
            v = row[0].value
            if isinstance(v, (datetime, date)):
                row[0].number_format = "yyyy-mm-dd"
    return ws


def write_report(res: Result, a: pd.DataFrame, b: pd.DataFrame, out: Path,
                 labels=("A", "B"), fuzzy=None, dayfirst: bool = True) -> Path:
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    key = res.key
    diff_pairs = {d["pair"] for d in res.diffs}
    fuzzy_n = sum(1 for m in res.matched if m["how"] != "key")
    rows = [
        ("Reconciliation report", "demo with synthetic sample data"),
        ("Generated", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Dataset A", f"{labels[0]} ({len(a)} rows)"),
        ("Dataset B", f"{labels[1]} ({len(b)} rows)"),
        ("Key column", key),
        ("Fuzzy matching", f"{fuzzy[0]} >= {fuzzy[1]} (token_sort_ratio, numbers must agree)"
         if fuzzy else "off"),
        ("Normalise text", "trim, collapse spaces, ignore case"),
        ("Normalise numbers", "1,250.00 = 1250 = 1250.0"),
        ("Normalise dates", "2026-02-03 = 3 Feb 2026 = "
         + ("03/02/2026 (day first)" if dayfirst else "02/03/2026 (month first)")),
        ("Normalise names", "ignore punctuation and Ltd / Limited / Inc / Co"),
        ("", ""),
        ("Matched pairs", len(res.matched)),
        ("  of which by key", len(res.matched) - fuzzy_n),
        ("  of which by fuzzy name", fuzzy_n),
        ("Matched, identical", len(res.matched) - len(diff_pairs)),
        ("Matched, with differences", len(diff_pairs)),
        ("Differing cells", len(res.diffs)),
        ("Only in A", len(res.only_a)),
        ("Only in B", len(res.only_b)),
    ]
    for side, keys in res.duplicates.items():
        rows.append((f"Duplicate keys in {side}", ", ".join(keys)))
    for r in rows:
        summary.append(list(r))
    summary["A1"].font = Font(bold=True, size=13)
    for r in range(2, len(rows) + 1):
        summary.cell(r, 1).font = Font(bold=True)
        summary.cell(r, 2).alignment = Alignment(horizontal="left", wrap_text=True, vertical="top")
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 70

    # Matched: A and B side by side, differing cells highlighted
    cols = [key] + res.columns
    header = ["match", "score"] + [f"{c} (A)" for c in cols] + [f"{c} (B)" for c in cols]
    body, marks = [], []
    for k, m in enumerate(res.matched):
        ra, rb = a.iloc[m["a"]], b.iloc[m["b"]]
        body.append([m["how"], m["score"]] + [ra[c] for c in cols] + [rb[c] for c in cols])
        marks.append({d["column"] for d in res.diffs if d["pair"] == k})
    ws = _sheet(wb, "Matched", header, body)
    for r, (m, cset) in enumerate(zip(res.matched, marks), start=2):
        if m["how"] != "key":
            for c in (1, 2):
                ws.cell(r, c).fill = FUZZY_FILL
        for c in cset:
            ci = cols.index(c)
            ws.cell(r, 3 + ci).fill = DIFF_FILL
            ws.cell(r, 3 + len(cols) + ci).fill = DIFF_FILL

    for title, df, idx in (("Only in A", a, res.only_a), ("Only in B", b, res.only_b)):
        ws = _sheet(wb, title, list(df.columns), [list(df.iloc[i]) for i in idx])
        for r in range(2, len(idx) + 2):
            ws.cell(r, 1).fill = ONLY_FILL

    drows = []
    for d in res.diffs:
        m = res.matched[d["pair"]]
        drows.append([a.iloc[m["a"]][key], b.iloc[m["b"]][key], m["how"], d["column"], d["a"], d["b"]])
    ws = _sheet(wb, "Value differences",
                [f"{key} (A)", f"{key} (B)", "match", "column", "value in A", "value in B"], drows)
    for r in range(2, len(drows) + 2):
        ws.cell(r, 5).fill = DIFF_FILL
        ws.cell(r, 6).fill = DIFF_FILL

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


# ------------------------------------------------------------------ CLI
def parse_fuzzy(text: str) -> tuple[str, int]:
    col, _, thr = text.rpartition(":")
    if not col or not thr.isdigit() or not 0 < int(thr) <= 100:
        raise argparse.ArgumentTypeError("use COLUMN:THRESHOLD, e.g. name:90")
    return col, int(thr)


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Reconcile two datasets by key")
    ap.add_argument("a", help="CSV, XLSX, or XLSX#Sheet")
    ap.add_argument("b", help="CSV, XLSX, or XLSX#Sheet")
    ap.add_argument("--key", required=True)
    ap.add_argument("--fuzzy", type=parse_fuzzy, metavar="COLUMN:THRESHOLD")
    ap.add_argument("--monthfirst", action="store_true", help="read 03/02/2026 as 2 March")
    ap.add_argument("-o", "--output", default="reconcile_report.xlsx")
    args = ap.parse_args(argv)
    a, la = load(args.a)
    b, lb = load(args.b)
    res = reconcile(a, b, args.key, args.fuzzy, dayfirst=not args.monthfirst)
    write_report(res, a, b, Path(args.output), (la, lb), args.fuzzy, not args.monthfirst)
    fz = sum(1 for m in res.matched if m["how"] != "key")
    print(f"matched {len(res.matched)} (fuzzy {fz}), only in A {len(res.only_a)}, "
          f"only in B {len(res.only_b)}, differing cells {len(res.diffs)} "
          f"in {len({d['pair'] for d in res.diffs})} rows -> {args.output}")
    for side, keys in res.duplicates.items():
        print(f"warning: duplicate keys in {side}: {', '.join(keys)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
