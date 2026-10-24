"""Render docs/*.png from samples/reconcile_report.xlsx.

LibreOffice is not required: cell values, fills, bold fonts and column
widths are read with openpyxl and drawn as a spreadsheet-style grid with
PyMuPDF. The picture is a faithful rendering of the workbook, not a mock-up.
"""
from datetime import date, datetime
from pathlib import Path
import sys

import fitz
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import make_samples  # noqa: E402
import reconcile as rc  # noqa: E402

ROW_H, CHAR_W, FS = 15, 5.2, 7.5


def rgb(hexstr):
    h = hexstr[-6:]
    return tuple(int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))


def fmt(v):
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float):
        return f"{v:,.2f}"
    return str(v)


def draw_sheet(page, ws, x0, y0, max_rows, title, all_tabs):
    widths = [max(ws.column_dimensions[c[0].column_letter].width or 8, 6) * CHAR_W
              for c in ws.iter_cols(max_row=1)]
    rows = list(ws.iter_rows(max_row=min(ws.max_row, max_rows)))
    total_w = 26 + sum(widths)
    page.insert_text((x0, y0 - 8), title, fontname="hebo", fontsize=10)
    y = y0
    # column letters
    page.draw_rect(fitz.Rect(x0, y, x0 + total_w, y + ROW_H), color=None, fill=(0.93, 0.93, 0.93))
    x = x0 + 26
    for i, w in enumerate(widths):
        page.insert_text((x + w / 2 - 3, y + 11), rows[0][i].column_letter, fontname="helv",
                         fontsize=FS, color=(0.4, 0.4, 0.4))
        x += w
    y += ROW_H
    for r in rows:
        page.draw_rect(fitz.Rect(x0, y, x0 + 26, y + ROW_H), color=None, fill=(0.93, 0.93, 0.93))
        page.insert_text((x0 + 6, y + 11), str(r[0].row), fontname="helv", fontsize=FS, color=(0.4, 0.4, 0.4))
        x = x0 + 26
        for c, w in zip(r, widths):
            rect = fitz.Rect(x, y, x + w, y + ROW_H)
            fill = (1, 1, 1)
            if c.fill is not None and c.fill.fill_type == "solid":
                fill = rgb(c.fill.fgColor.rgb)
            page.draw_rect(rect, color=(0.82, 0.82, 0.82), fill=fill, width=0.4)
            text = fmt(c.value)
            bold = bool(c.font and c.font.b)
            color = rgb(c.font.color.rgb) if (c.font and c.font.color and isinstance(c.font.color.rgb, str)) else (0, 0, 0)
            font = "hebo" if bold else "helv"
            maxw = w - 6
            while text and fitz.get_text_length(text, fontname=font, fontsize=FS) > maxw:
                text = text[:-1]
            right = (isinstance(c.value, (int, float)) and not isinstance(c.value, bool)
                     and not (c.alignment and c.alignment.horizontal == "left"))
            tx = x + w - 3 - fitz.get_text_length(text, fontname=font, fontsize=FS) if right else x + 3
            page.insert_text((tx, y + 11), text, fontname=font, fontsize=FS, color=color)
            x += w
        y += ROW_H
    # sheet tabs
    tx = x0
    for name in all_tabs:
        w = fitz.get_text_length(name, fontname="helv", fontsize=FS) + 16
        active = name == ws.title
        page.draw_rect(fitz.Rect(tx, y + 4, tx + w, y + 20), color=(0.7, 0.7, 0.7),
                       fill=(1, 1, 1) if active else (0.9, 0.9, 0.9), width=0.5)
        page.insert_text((tx + 8, y + 15), name, fontname="hebo" if active else "helv",
                         fontsize=FS, color=(0.1, 0.4, 0.2) if active else (0.3, 0.3, 0.3))
        tx += w + 2
    return total_w, y + 24 - y0


def legend(page, x, y):
    for fill, label in (("FFE08A", "value differs (after normalisation)"),
                        ("D6E9FF", "paired by fuzzy name"), ("F8D7DA", "row only on one side")):
        page.draw_rect(fitz.Rect(x, y - 8, x + 12, y + 2), color=(0.6, 0.6, 0.6), fill=rgb(fill), width=0.4)
        page.insert_text((x + 16, y), label, fontname="helv", fontsize=8)
        x += 30 + fitz.get_text_length(label, fontname="helv", fontsize=8)


def render(report: Path, out: Path, sheets, title):
    wb = load_workbook(report)
    tabs = wb.sheetnames
    # measure first
    tmp = fitz.open()
    p = tmp.new_page(width=4000, height=4000)
    sizes = [draw_sheet(p, wb[s], 0, 30, n, s, tabs) for s, n, _ in sheets]
    tmp.close()
    pad = 20
    width = max(w for w, _ in sizes) + 2 * pad
    height = 50 + sum(h + 40 for _, h in sizes) + 10
    doc = fitz.open()
    page = doc.new_page(width=width, height=height)
    page.draw_rect(page.rect, color=None, fill=(0.98, 0.98, 0.97))
    page.insert_text((pad, 24), title, fontname="hebo", fontsize=13)
    legend(page, pad, 42)
    y = 76
    for (s, n, caption), (_, h) in zip(sheets, sizes):
        draw_sheet(page, wb[s], pad, y, n, caption, tabs)
        y += h + 40
    page.get_pixmap(dpi=170).save(out)
    doc.close()


if __name__ == "__main__":
    paths = make_samples.build(ROOT / "samples")
    report = ROOT / "samples" / "reconcile_report.xlsx"
    rc.main([str(paths["a_xlsx"]), str(paths["b_xlsx"]), "--key", "id", "--fuzzy", "name:90",
             "-o", str(report)])
    docs = ROOT / "docs"
    docs.mkdir(exist_ok=True)
    render(report, docs / "report_matched.png",
           [("Matched", 31, "Sheet 'Matched': A and B side by side (all 28 matched pairs)")],
           "reconcile_report.xlsx: matched rows, differing cells highlighted")
    render(report, docs / "report_summary_diffs.png",
           [("Summary", 20, "Sheet 'Summary'"), ("Value differences", 20, "Sheet 'Value differences'"),
            ("Only in B", 10, "Sheet 'Only in B' (the decoy 'Sample Co 071' was not fuzzy-matched)")],
           "reconcile_report.xlsx: summary, value differences, only-in-B")
    print("ok")
